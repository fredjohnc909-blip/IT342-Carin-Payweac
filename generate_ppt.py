import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()

    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "PayWEAC - Rent Payment & Tracking Platform"
    subtitle.text = "IT342 - Systems Integration and Architecture\nFinal Project Presentation\nWilfred John R. Carin"

    # Slide 2: Self-Introduction
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.shapes.placeholders[1]
    title.text = "Self-Introduction"
    tf = body.text_frame
    tf.text = "Name: Wilfred John R. Carin"
    p = tf.add_paragraph()
    p.text = "Course/Section: IT342 - Systems Integration and Architecture"
    p = tf.add_paragraph()
    p.text = "Project Focus: Simplifying rent tracking and payments through a centralized digital platform."

    # Slide 3: System Project Introduction
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.shapes.placeholders[1]
    title.text = "System Project Introduction"
    tf = body.text_frame
    tf.text = "Purpose: A digital platform for tenants to track rent dues and submit payments."
    p = tf.add_paragraph()
    p.text = "Problem Addressed: Manual tracking of rent payments, lost receipts, and unorganized rent histories."
    p = tf.add_paragraph()
    p.text = "Intended Users: Tenants (for payments/tracking) and Administrators (for managing/approving payments)."
    p = tf.add_paragraph()
    p.text = "Overall Goal: To streamline the rent collection process, making it transparent, efficient, and easy to use."

    # Slide 4: Main Features of the System
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.shapes.placeholders[1]
    title.text = "Main Features of PayWEAC"
    tf = body.text_frame
    tf.text = "User Authentication and Authorization (JWT)"
    p = tf.add_paragraph()
    p.text = "Role-Based Access Control (Admin vs Tenant views)"
    p = tf.add_paragraph()
    p.text = "CRUD Operations for Rents and Payments"
    p = tf.add_paragraph()
    p.text = "File Uploads (Uploading payment receipts)"
    p = tf.add_paragraph()
    p.text = "Payment Approval Workflow (Admin verification of submitted payments)"
    
    # Slide 5: System Architecture
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.shapes.placeholders[1]
    title.text = "System Architecture"
    tf = body.text_frame
    tf.text = "Client-Server Architecture"
    p = tf.add_paragraph()
    p.text = "Frontend: React 18 (Vite)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Backend: Java 17, Spring Boot 3"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Database: MySQL"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "REST API Architecture (JSON communication)"
    p = tf.add_paragraph()
    p.text = "Layered Backend Architecture (Controllers, Services, Repositories, Entities)"

    # Slide 6: Component Interaction & Data Flow
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.shapes.placeholders[1]
    title.text = "Component Interaction & Data Flow"
    tf = body.text_frame
    tf.text = "1. User initiates a request on the React Frontend."
    p = tf.add_paragraph()
    p.text = "2. Frontend sends an HTTP REST API request with JWT to Spring Boot Backend."
    p = tf.add_paragraph()
    p.text = "3. Spring Security validates the JWT token."
    p = tf.add_paragraph()
    p.text = "4. The RestController receives the request and processes logic."
    p = tf.add_paragraph()
    p.text = "5. The Repository interfaces with the MySQL Database via Spring Data JPA."
    p = tf.add_paragraph()
    p.text = "6. Backend sends JSON response back to Frontend."

    # Slide 7: Proof of Implementation
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.shapes.placeholders[1]
    title.text = "Proof of Implementation"
    tf = body.text_frame
    tf.text = "Separation of Concerns: Clear /web and /backend directories."
    p = tf.add_paragraph()
    p.text = "REST Controllers: RentController, PaymentController, AdminController."
    p = tf.add_paragraph()
    p.text = "Database Entities: Mapping User, Rent, and Payment tables."
    p = tf.add_paragraph()
    p.text = "Frontend API Calls: Using fetch/axios to communicate with backend endpoints."

    # Slide 8: System Demonstration
    title_only_slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(title_only_slide_layout)
    title = slide.shapes.title
    title.text = "System Demonstration"
    
    left = Inches(2)
    top = Inches(2.5)
    width = Inches(6)
    height = Inches(1)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "(Switch to Live Demo or Screen Recording)"
    p.font.size = Pt(32)
    p.font.italic = True
    p.alignment = PP_ALIGN.CENTER

    prs.save('PayWEAC_Final_Presentation.pptx')

if __name__ == '__main__':
    create_presentation()
